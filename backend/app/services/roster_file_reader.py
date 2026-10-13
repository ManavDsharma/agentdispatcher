"""Reads an uploaded .xlsx or .csv into plain rows of text.

Only the two agreed formats are accepted, and the content is checked, not just
the filename (an .exe renamed to .xlsx is rejected). No interpretation of the
roster happens here — see app/domain/roster_schema.py for that.
"""

from __future__ import annotations

import csv
import io
import os
from dataclasses import dataclass
from datetime import date, datetime, time
from typing import List

MAX_FILE_BYTES = 5 * 1024 * 1024
MAX_ROWS = 5000  # data rows, header excluded
MAX_COLUMNS = 40
MAX_SHEETS = 10
ALLOWED_EXTENSIONS = (".xlsx", ".csv")


class RosterFileError(Exception):
    """A problem with the file itself (type, size, unreadable), not its contents."""

    def __init__(self, message: str, status_code: int = 400):
        super().__init__(message)
        self.status_code = status_code


@dataclass
class Sheet:
    name: str  # sheet name for .xlsx, the filename for .csv
    rows: List[List[str]]  # row i is spreadsheet row i + 1; blank rows are kept so numbers line up


def _cell_text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "Yes" if value else "No"
    if isinstance(value, float):
        # Excel stores 2026 as 2026.0
        return str(int(value)) if value.is_integer() else str(value)
    if isinstance(value, datetime):
        return value.isoformat(sep=" ", timespec="seconds") if (value.hour or value.minute or value.second) else value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, time):
        return value.strftime("%H:%M")
    return str(value).strip()


def _trim_trailing_blank(rows: List[List[str]]) -> List[List[str]]:
    while rows and not any(cell for cell in rows[-1]):
        rows.pop()
    return rows


def _read_xlsx(content: bytes) -> List[Sheet]:
    if not content.startswith(b"PK\x03\x04"):
        raise RosterFileError("This isn't a valid .xlsx workbook. Save it as 'Excel Workbook (.xlsx)' and try again.")

    try:
        from openpyxl import load_workbook

        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception:
        raise RosterFileError("Couldn't open the workbook. It may be corrupted or password-protected.")

    sheets: List[Sheet] = []
    try:
        for worksheet in workbook.worksheets[:MAX_SHEETS]:
            if worksheet.sheet_state != "visible":
                continue
            rows: List[List[str]] = []
            for raw in worksheet.iter_rows(max_col=MAX_COLUMNS, values_only=True):
                rows.append([_cell_text(v) for v in raw])
                if len(rows) > MAX_ROWS + 20:  # +20 leaves room for title rows above the header
                    raise RosterFileError(
                        f"The sheet '{worksheet.title}' has more than {MAX_ROWS} rows. Split it into smaller uploads.",
                        413,
                    )
            sheets.append(Sheet(worksheet.title, _trim_trailing_blank(rows)))
    except RosterFileError:
        raise
    except Exception:
        raise RosterFileError("Couldn't read the workbook contents. It may be corrupted.")
    finally:
        workbook.close()

    return sheets


def _read_csv(content: bytes, filename: str) -> List[Sheet]:
    if b"\x00" in content[:8192]:
        raise RosterFileError("This doesn't look like a CSV text file.")

    try:
        text = content.decode("utf-8-sig")
    except UnicodeDecodeError:
        text = content.decode("cp1252", errors="replace")  # Excel's default "CSV" on Windows

    sample = text[:4096]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t")
    except csv.Error:
        dialect = csv.excel

    rows: List[List[str]] = []
    try:
        for raw in csv.reader(io.StringIO(text), dialect):
            rows.append([c.strip() for c in raw[:MAX_COLUMNS]])
            if len(rows) > MAX_ROWS + 20:
                raise RosterFileError(f"The file has more than {MAX_ROWS} rows. Split it into smaller uploads.", 413)
    except csv.Error:
        raise RosterFileError("Couldn't parse the CSV file.")

    return [Sheet(filename, _trim_trailing_blank(rows))]


def read_upload(filename: str, content: bytes) -> List[Sheet]:
    """Parse the upload into sheets of text rows, or raise RosterFileError."""
    extension = os.path.splitext(filename or "")[1].lower()
    if extension not in ALLOWED_EXTENSIONS:
        raise RosterFileError("Only .xlsx and .csv files can be uploaded.", 415)
    if not content:
        raise RosterFileError("The file is empty.")
    if len(content) > MAX_FILE_BYTES:
        raise RosterFileError(f"The file is larger than {MAX_FILE_BYTES // (1024 * 1024)} MB.", 413)

    sheets = _read_xlsx(content) if extension == ".xlsx" else _read_csv(content, filename)
    if not any(sheet.rows for sheet in sheets):
        raise RosterFileError("The file has no data.")
    return sheets
